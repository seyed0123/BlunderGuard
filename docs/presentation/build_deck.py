"""Build the 15-minute deck. Requires python-pptx and Pillow; run from any cwd."""
from pathlib import Path
import json, csv
from decimal import Decimal, ROUND_HALF_UP
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
from PIL import Image
ROOT=Path(__file__).resolve().parents[2]; OUT=Path(__file__).resolve().parent
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
NAVY='10283F'; TEAL='007F86'; GREEN='86B94F'; WHITE='FFFFFF'; GRAY='526576'; LIGHT='EAF3F5'; RED='AC493E'
rows=list(csv.DictReader(open(ROOT/'docs/report/chapter3_assets/verified_results.csv')))
NAMES={'gpt-5.5':'GPT-5.5','qwen3.7-plus':'Qwen3.7 Plus','gemini-3.5-flash':'Gemini 3.5 Flash','grok-4.3':'Grok 4.3','cf.llama-3.3-70b-instruct-fp8-fast':'Llama 3.3 70B (FP8)','gpt-oss-120b':'GPT-OSS 120B','qwen3-30b-a3b':'Qwen3 30B A3B','gemma-3-12b-it-Q4_0':'Gemma 3 12B (Q4)','cf.gemma-4-26b-a4b-it':'Gemma 4 26B A4B','Llama-3.2-3B-Instruct-Q4_K_M':'Llama 3.2 3B (Q4)','chess-q8_0':'Chess Q8 (Qwen3 1.7B)'}
short=[NAMES[r['model']] for r in rows]
by_model={r['model']:r for r in rows}
local=by_model['chess-q8_0']
def fmt(v): return str(Decimal(str(v)).quantize(Decimal('.01'),rounding=ROUND_HALF_UP))
def rect(s,x,y,w,h,c):
 a=s.shapes.add_shape(1,Inches(x),Inches(y),Inches(w),Inches(h));a.fill.solid();a.fill.fore_color.rgb=RGBColor.from_string(c);a.line.fill.background();return a
def text(s,t,x,y,w,h,size=24,color=NAVY,bold=False,font='Arial',rtl=False):
 a=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=a.text_frame;tf.word_wrap=True
 tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(t.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name=font;p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=RGBColor.from_string(color);p.space_after=Pt(12)
  if rtl:p._p.get_or_add_pPr().set('rtl','1')
 return a
notes=[]
def slide(title,kicker='BLUNDERGUARD',backup=False):
 s=prs.slides.add_slide(prs.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(WHITE)
 rect(s,0,0,13.333,.13,TEAL);text(s,kicker,.6,.35,12,.3,12,TEAL,True);text(s,title,.6,.92,12.1,.95,31,NAVY,True)
 n=len(prs.slides);text(s,'Shahid Beheshti University',.6,7.12,10,.2,10,GRAY);text(s,f'{n}/12' if not backup else 'BACKUP',11.5,7.08,1.2,.3,12,GRAY)
 if backup:s._element.set('show','0')
 return s
def note(s,sec,t):
 title=next(sh.text for sh in s.shapes if sh.has_text_frame and sh.top==Inches(.92))
 content=f'زمان پیشنهادی: {sec} ثانیه\n\n{t}' if sec else 'اسلاید پشتیبان؛ فقط در صورت سؤال نمایش داده شود.\n\n'+t
 tf=s.notes_slide.notes_text_frame;tf.text=content
 for p in tf.paragraphs:
  p.font.name='B Nazanin';p.font.size=Pt(14);p._p.get_or_add_pPr().set('rtl','1')
 notes.append((title,sec,t))
def pic(s,path,x,y,w,h):
 path=Path(path);iw,ih=Image.open(path).size;scale=min(w/iw,h/ih);s.shapes.add_picture(str(path),Inches(x+(w-iw*scale)/2),Inches(y+(h-ih*scale)/2),width=Inches(iw*scale),height=Inches(ih*scale))
def call(s,t,y=6.35):rect(s,.6,y,12.1,.52,LIGHT);text(s,t,.8,y+.09,11.7,.4,18,TEAL,True)
def bullets(s,items,x=.8,y=2,w=11.7,size=25,gap=1):
 for i,t in enumerate(items):text(s,t,x,y+i*gap,w,gap-.1,size)
# 1
s=slide('Building a Chess Analysis Model','EXPERT SYSTEMS + LARGE LANGUAGE MODELS')
text(s,'به نام خدا',5.4,.32,2.5,.45,20,TEAL,False,'B Nazanin',True)
pic(s,ROOT/'docs/report/chapter3_assets/beheshti_logo.png',11.45,.2,1.05,.75)
text(s,'BlunderGuard',.75,2.15,8,1,45,TEAL,True);text(s,'From engine evaluations\nto learner-facing commentary',.75,3.2,8,1.3,30)
text(s,'Seyed Ali Hosseini Nasab\nSupervisor: Dr. Mehrnoush Shamsfard',.75,5.45,8,1,21,GRAY)
pic(s,OUT/'assets/slide1.jpeg',9,2.65,3.3,3.6)
note(s,30,'سلام. موضوع پروژه من تولید توضیح قابل فهم برای حرکت‌های شطرنج با ترکیب موتور تحلیل و مدل زبانی است. نام پیاده‌سازی BlunderGuard است. تصویر اولیه ارائه را نگه داشته‌ام تا مسئله را یادآوری کند: دیدن اشتباه آسان‌تر از فهمیدن علت آن است. در این ارائه ابتدا روش، سپس آزمایش و در پایان محدودیت‌های نتایج را توضیح می‌دهم.')
# 2
s=slide('A score tells us how much—not why')
pic(s,OUT/'assets/slide9.png',.65,2,7.3,3.75)
bullets(s,['Engine: evaluation + best line','Learner: threat + consequence','Goal: explain the move'],8.3,2.25,4.3,24,1.12)
call(s,'The challenge is translating chess evidence into a useful explanation.')
note(s,60,'در تصویر، خروجی عددی و دنباله حرکت‌های یک موتور شطرنج دیده می‌شود. بازیکن باتجربه ممکن است از روی این خروجی علت خطا را پیدا کند، ولی بازیکن مبتدی معمولاً نمی‌داند چه تهدیدی ایجاد شده و چرا حرکت بهتری وجود داشته است. هدف پروژه جایگزین کردن موتور نیست؛ موتور ارزیابی می‌کند و مدل زبانی باید همان شواهد را توضیح دهد. مثال ساده این است که به جای گفتن امتیاز کم شد، توضیح دهیم کدام مهره بی‌دفاع مانده یا چه تهدیدی نادیده گرفته شده است. این تصویر از ارائه اولیه حفظ شده و صرفاً شکل خروجی موتور را نشان می‌دهد. در ادامه سؤال اصلی این است که آیا متن تولیدشده واقعاً با شواهد سازگار می‌ماند؟')
# 3
s=slide('Can an LLM explain engine evidence faithfully?')
for x,num,title,body in [( .7,'01','Build','A local analysis pipeline\nwith engine + language model'),(4.9,'02','Adapt','A small model trained\non chess-domain text'),(9.1,'03','Evaluate','11 models on the same\n60 benchmark positions')]:
 rect(s,x,2.1,3.55,3.65,LIGHT);text(s,num,x+.2,2.3,3,.6,33,TEAL,True);text(s,title,x+.2,3.15,3,.5,27,NAVY,True);text(s,body,x+.2,4,3.1,1.5,23)
call(s,'Success requires chess correctness and faithfulness—not just fluent prose.')
note(s,60,'سه بخش اصلی کار را از هم جدا می‌کنم. بخش اول ساخت سامانه‌ای است که موقعیت و حرکت را به موتور می‌دهد و شواهد را برای توضیح آماده می‌کند. بخش دوم سازگار کردن یک مدل کوچک با متن‌های حوزه شطرنج است. بخش سوم ارزیابی خروجی مدل‌ها روی موقعیت‌های مشترک است. این تفکیک مهم است، چون کار کردن نرم‌افزار به معنی درست بودن توضیح نیست و آموزش روی متن شطرنج نیز به تنهایی درک صفحه را ثابت نمی‌کند. پرسش پژوهشی من درباره وفاداری توضیح به شواهد است. برای پاسخ، هم امتیاز کلی را می‌بینیم و هم مؤلفه‌های صحت و وفاداری را جداگانه بررسی می‌کنیم.')
# 4
s=slide('Engine analysis becomes structured context')
pic(s,OUT/'assets/slide20.png',.7,1.85,7.3,4.4)
bullets(s,['Input: position + played move','Stockfish: depth 17, up to 3 lines','Before/after feature differences','LLM: short learner-facing text'],8.2,2,4.4,22,1.02)
call(s,'Engine evidence guides generation; the output still needs verification.')
note(s,90,'این شکل همان نمودار مسیر کلی در ارائه اولیه است. ورودی، وضعیت صفحه و حرکت انجام‌شده است. در پیاده‌سازی، وضعیت قبل و بعد از حرکت تحلیل می‌شود؛ عمق پیش‌فرض استاک‌فیش ۱۷ است و تا سه ادامه بررسی می‌شود. علاوه بر خروجی موتور، ویژگی‌هایی مانند مهره‌های آویزان، امنیت شاه، مرکز و ساختار پیاده با قواعد برنامه استخراج می‌شوند. تفاوت ویژگی‌های قبل و بعد به ساخت زمینه برای مدل کمک می‌کند. مدل سپس توضیح کوتاهی برای کاربر تولید می‌کند. باید تأکید کنم که این قواعد بی‌خطا نیستند و قرار دادن شواهد در ورودی نیز تضمین نمی‌کند مدل از آن‌ها درست استفاده کند. نمودار، مسیر طراحی سامانه را نشان می‌دهد؛ اعتبار توضیح را باید جداگانه با ارزیابی سنجید. در نسخه محلی، مدل از طریق llama.cpp اجرا می‌شود و رابط کاربر به سرویس پشتی متصل است.')
# 5
s=slide('A small local model: what was actually trained')
bullets(s,['Base model: Qwen3-1.7B','LoRA adaptation on unpaired chess-book text','Local inference: Q8_0 GGUF with llama.cpp'],.8,2,11.5,28,1.04)
rect(s,.7,5.3,11.9,.75,LIGHT);text(s,'Domain vocabulary ≠ supervised position-to-explanation learning',.9,5.47,11.5,.45,24,TEAL,True)
note(s,75,'مدل پایه در دفترچه آموزش Qwen3 با اندازه ۱٫۷ میلیارد پارامتر است. برای کاهش هزینه آموزش از LoRA استفاده شده و داده آموزشی متن کتاب‌های شطرنج بوده است. نکته تعیین‌کننده این است که این متن‌ها جفت‌های صریح موقعیت صفحه و توضیح صحیح نیستند. بنابراین انتظار داریم مدل با زبان و اصطلاحات شطرنج آشنا شود، ولی از این آموزش به تنهایی نمی‌توان یادگیری تحلیل دقیق یک موقعیت را نتیجه گرفت. خروجی برای اجرای محلی به قالب GGUF با کوانتیزه‌سازی Q8 تبدیل شده است. در نتایج این مدل با نام Chess Q8 دیده می‌شود. چون مقایسه کنترل‌شده قبل و بعد از آموزش نداریم، ادعا نمی‌کنم که آموزش حتماً کیفیت تحلیل را بهتر کرده است. این محدودیت را هنگام تفسیر نتایج نگه می‌داریم.')
# 6
s=slide('One shared benchmark, 660 judgments')
for x,n,label in [(.8,'60','positions'),(4.95,'11','candidate models'),(9.1,'660','scored responses')]:
 text(s,n,x,2,3.5,.95,54,TEAL,True);text(s,label,x,3.02,3.5,.5,25)
bullets(s,['20 positions per rating group: beginner / intermediate / advanced','Candidates aligned by group and row ID','One fixed judge: gemma4-12b; same reference for each position'],.8,4.15,11.8,23,.66)
call(s,'Rating groups describe source games; they are not measured difficulty levels.')
note(s,75,'مجموعه ارزیابی شامل ۶۰ موقعیت است: ۲۰ موقعیت از هر گروه مبتدی، متوسط و پیشرفته. گروه‌بندی با میانگین ریتینگ دو بازیکن بازی مبدأ انجام شده؛ کمتر از ۱۲۰۰، از ۱۲۰۰ تا کمتر از ۱۸۰۰، و ۱۸۰۰ به بالا. پس این برچسب‌ها را نباید اندازه‌گیری مستقیم دشواری موقعیت دانست. برای هر موقعیت خروجی ۱۱ مدل وجود دارد و در مجموع ۶۶۰ پاسخ امتیاز گرفته‌اند. تطبیق با گروه و شناسه ردیف انجام شده تا مقایسه بر مبنای موقعیت مشترک باشد. داور ثابت gemma4-12b است. در عین حال، اجرای مدل‌ها الزاماً سخت‌افزار یا تنظیمات تولید یکسان ندارد؛ این آزمایش مقایسه پاسخ‌های ثبت‌شده است، نه آزمایش سرعت یا هزینه.')
# 7
s=slide('The score rewards correctness and faithfulness')
for i,(label,weight) in enumerate([('Correctness',35),('Faithfulness',25),('Helpfulness',20),('Instruction following',10),('Fluency',10)]):
 y=1.9+i*.69;text(s,label,.8,y,4,.45,24);rect(s,5,y+.04,weight*.16,.36,TEAL if i<2 else '79B9BF');text(s,f'{weight}%',11,y,1,.45,24,TEAL,True)
call(s,'The judge compares candidate and reference text; it does not inspect the board.')
note(s,75,'هر پاسخ در پنج مؤلفه از صفر تا ده امتیاز می‌گیرد. صحت ۳۵ درصد و وفاداری ۲۵ درصد وزن دارند؛ در نتیجه مجموعاً ۶۰ درصد امتیاز به این دو اختصاص یافته است. مفید بودن ۲۰ درصد و پیروی از دستور و روانی هر کدام ۱۰ درصد هستند. امتیاز کلی میانگین وزنی این مؤلفه‌هاست. محدودیت مهم اینجاست که داور متن پاسخ و متن مرجع را مقایسه می‌کند و مستقیماً صفحه را با موتور بررسی نمی‌کند. بنابراین بهتر است نتایج را هم‌سویی با مرجع تحت این داور بنامیم و نه اثبات مستقل صحت تمام جزئیات شطرنجی. همچنین روانی بالا می‌تواند بخشی از امتیاز را بالا ببرد حتی وقتی محتوای شطرنجی ضعیف است. در اسلاید بعد رتبه کلی را می‌بینیم و سپس این اختلاف را باز می‌کنیم.')
# 8
s=slide('Higher overall scores do not imply perfect analysis')
for i,(r,label) in enumerate(zip(rows,short)):
 y=1.87+i*.365;v=float(r['overall']);text(s,label,.8,y,4.2,.32,17);rect(s,5,y+.04,6*v/10,.22,TEAL if i<3 else (RED if r['model']=='chess-q8_0' else '83A8BA'));text(s,fmt(r['overall']),11.35,y,1,.35,18,NAVY,True)
text(s,'Mean weighted score / 10  •  60 responses per model',5,6.03,7,.35,15,GRAY)
call(s,'Best observed: GPT-5.5 7.31/10  |  Chess Q8 (1.7B): 2.86/10')
note(s,90,'این نمودار نتایج نهایی ۱۱ مدل را نشان می‌دهد. بهترین میانگین همچنان GPT-5.5 با ۷٫۳۱ است. امتیاز مدل محلی مبتنی بر Qwen3-1.7B اکنون ۲٫۸۶ است. نکته جالب، نزدیکی مدل ۱٫۷ میلیاردپارامتری به چند مدل بزرگ‌تر است. در این اسلاید تمام ۱۱ مدل آمده‌اند تا مقایسه انتخابی تصویر کلی را پنهان نکند. فاصله با مدل‌های قوی‌تر همچنان قابل توجه است. اسلاید بعد مقایسه نزدیک‌تر مدل کوچک را نشان می‌دهد.')
# 9
s=slide('A 1.7B model scores close to several larger models')
text(s,'Qwen3-1.7B',.8,1.95,4.6,.5,27,TEAL,True)
text(s,'2.86 / 10',.8,2.75,4.6,.9,42,NAVY,True)
text(s,'1.7 billion parameters\nLocal GGUF inference\nChess-domain adaptation',.8,3.85,4.5,1.7,23)
comparisons=[('chess-q8_0','Chess Q8 · 1.7B'),('Llama-3.2-3B-Instruct-Q4_K_M','Llama 3.2 · 3B'),('cf.gemma-4-26b-a4b-it','Gemma 4 · 26B A4B'),('gemma-3-12b-it-Q4_0','Gemma 3 · 12B')]
for i,(model,label) in enumerate(comparisons):
 y=2.02+i*.85;v=float(by_model[model]['overall']);c=TEAL if i==0 else '83A8BA'
 text(s,label,5.55,y,5.8,.36,21,c,True);rect(s,5.55,y+.44,5.8*v/10,.19,c);text(s,fmt(by_model[model]['overall']),11.55,y+.3,.9,.4,22,c,True)
text(s,'Mean score / 10 · same 60 positions and judge',5.55,5.7,6.8,.35,16,GRAY)
call(s,'Similar benchmark means; general capability and efficiency are not established.')
note(s,75,'مدل اختصاصی با ۱٫۷ میلیارد پارامتر به میانگین ۲٫۸۵۸۳ رسیده است. مدل Llama با سه میلیارد پارامتر ۲٫۸۵۵۰ دارد، پس هر دو با دو رقم اعشار ۲٫۸۶ دیده می‌شوند. این اختلاف بسیار کوچک را برتری معنادار نمی‌نامیم. امتیاز مدل ما به Gemma 4 با ۲٫۹۰ و Gemma 3 با ۲٫۹۵ هم نزدیک است. این نتیجه برای یک مدل کوچک امیدوارکننده است و نشان می‌دهد اندازه بزرگ‌تر در این آزمایش به تنهایی تضمین امتیاز بهتر نیست. بااین‌حال، معماری و آموزش مدل‌ها متفاوت است و مدل 26B A4B تعداد پارامتر فعال متفاوتی از تعداد کل دارد. سرعت یا مصرف منابع را اندازه نگرفته‌ایم. نتیجه دقیق، نزدیکی میانگین در همین مجموعه است.')
# 10
s=slide('A real failure: the explanation names the wrong move')
p=ROOT/'artifacts/analysis_images/g1_m38__model_chess-q8_0.png'
pic(s,p,.6,1.82,7.55,5.1)
text(s,'Played move: Bxf8',8.5,1.95,4.2,.5,25,TEAL,True)
text(s,'Model output (excerpt)',8.5,2.55,4.2,.4,20,GRAY,True)
text(s,'“White\'s move Nxe4 weakens Black\'s position…”',8.5,3.15,4.2,1.4,26,RED,True)
text(s,'Correctness  0/10\nFaithfulness  0/10\nFluency       10/10',8.5,4.85,4.2,1.5,23)
text(s,'advanced / g1_m38 · overall 2.00/10',8.5,6.5,4.2,.4,15,GRAY)
note(s,105,'این نمونه مستقیماً از ردیف advanced با شناسه g1_m38 در مجموعه ارزیابی آمده است. حرکت ثبت‌شده Bxf8 است؛ فیل رخ را در خانه f8 می‌گیرد. ابتدا چند ثانیه به دو صفحه نگاه کنید تا تغییر موقعیت روشن شود. مدل محلی در شروع پاسخ به جای این حرکت درباره Nxe4 صحبت می‌کند و سپس حرکت و تهدیدهای دیگری را توضیح می‌دهد. همین عدم تطبیق حرکت نشان می‌دهد پاسخ به موقعیت مورد نظر متصل نمانده است. طبق داور، صحت و وفاداری هر دو صفر شده‌اند، ولی روانی ده است. پیروی از دستور هم ده و مفید بودن صفر بوده، بنابراین امتیاز وزنی دو به دست آمده است. متن مرجع درباره بی‌دفاع ماندن فیل و پیامد تاکتیکی معامله هشدار می‌دهد؛ آن را به عنوان توضیح مرجع گزارش می‌کنم، نه تحلیل مستقلی که اینجا دوباره با موتور تأیید شده باشد. این نمونه برای نمایش یک شکست انتخاب شده و جایگزین میانگین کل داده‌ها نیست. پیام اصلی این است که کنترل نام حرکت و اتصال متن به صفحه باید بخشی از سامانه نهایی باشد.')
# 11
s=slide('From this prototype to a more reliable chess coach')
for x,n,title,body in [(.7,'01','Ground the training','Pair board positions with\nverified explanations'),(4.9,'02','Verify the claims','Check moves and threats\nagainst chess tools'),(9.1,'03','Test with people','Expand the benchmark\nand add expert review')]:
 rect(s,x,2.05,3.55,3.8,LIGHT)
 text(s,n,x+.2,2.25,3.1,.6,33,TEAL,True)
 text(s,title,x+.2,3.08,3.1,.85,25,NAVY,True)
 text(s,body,x+.2,4.12,3.1,1.45,23)
call(s,'Next experiment: compare base and adapted Qwen3-1.7B under identical settings.')
note(s,75,'نمونه قبل نشان داد که چالش اصلی، اتصال توضیح به وضعیت واقعی صفحه است. برای ادامه کار سه اولویت دارم. اول، به جای تکیه بر متن کتاب، جفت‌های موقعیت و توضیح تأییدشده تهیه شود. دوم، نام حرکت‌ها و ادعاهای تاکتیکی پیش از نمایش با ابزارهای شطرنج کنترل شوند. سوم، مجموعه آزمون گسترش پیدا کند و بازیکنان خبره کیفیت آموزشی متن را ارزیابی کنند. یک آزمایش مشخص هم مقایسه مدل پایه Qwen3-1.7B و نسخه سازگارشده در شرایط یکسان است تا سهم آموزش روشن شود. این موارد مسیر پیشنهادی ادامه کار هستند، نه قابلیت‌های تکمیل‌شده سامانه فعلی.')
# 12
s=slide('A small model, a practical starting point')
text(s,'1.7B',.8,2.0,4,.95,54,TEAL,True)
text(s,'parameters · Qwen3',.8,3.04,4.4,.5,24)
text(s,'2.86 / 10',.8,4.0,4.4,.85,40,NAVY,True)
text(s,'mean benchmark score',.8,4.96,4.4,.5,22,GRAY)
text(s,'What this project contributes',5.7,2.1,6.5,.6,28,NAVY,True)
bullets(s,['An engine-guided commentary pipeline','Scores close to selected larger models','A clear path to improve factual accuracy'],5.7,3.05,6.6,24,.86)
text(s,'Thank you',.8,6.15,5,.55,30,TEAL,True)
text(s,'Questions & discussion',5.7,6.25,6.6,.45,24,GRAY)
note(s,45,'در پایان، دستاورد پروژه یک سامانه قابل اجرا برای ترکیب تحلیل موتور و توضیح زبانی و یک چارچوب مشترک برای ارزیابی مدل‌هاست. مدل ۱٫۷ میلیاردپارامتری ما امتیاز ۲٫۸۶ گرفته که در همین آزمون به چند مدل بزرگ‌تر نزدیک است. این نتیجه انگیزه ادامه کار با مدل‌های سبک را فراهم می‌کند. در عین حال، خطاهای محتوایی باقی‌اند و برای کاربرد آموزشی قابل اعتماد باید اتصال متن به شواهد بهتر شود. پیام نهایی این است که مدل کوچک نقطه شروع عملی مناسبی برای ادامه پژوهش است؛ نه اینکه توانایی آن در همه کارها با مدل‌های بزرگ برابر باشد. از توجه شما ممنونم.')
# backups
s=slide('Full metric breakdown','BACKUP / RESULTS',True)
headers=['Model','Correct','Faithful','Helpful','Instr.','Fluent','Overall'];xs=[.75,5.1,6.3,7.5,8.7,9.9,11.1]
for x,h in zip(xs,headers):text(s,h,x,1.9,4.2 if x<1 else 1.1,.35,16,TEAL,True)
for i,(r,label) in enumerate(zip(rows,short)):
 y=2.4+i*.35
 if i%2==0:rect(s,.65,y-.025,12,.34,LIGHT)
 for j,value in enumerate([label]+[fmt(r[k]) for k in ['correctness','faithfulness','helpfulness','instruction_following','fluency','overall']]):text(s,value,xs[j],y,4.2 if j==0 else 1.1,.32,15)
note(s,0,'تمام مقدارها میانگین ۶۰ پاسخ هستند. نام کامل مدل‌ها و فایل‌های ورودی در فایل راهنمای ارائه ثبت شده‌اند. این اسلاید برای پاسخ به سؤال درباره مؤلفه‌های امتیاز است.')
s=slide('Training and implementation details','BACKUP / METHOD',True)
bullets(s,['LoRA: rank 16, alpha 32, dropout 0.05; q_proj and v_proj','Training: 1 epoch; learning rate 2e-4; text length 256','Batch size 2 × gradient accumulation 16','React + Flask; Stockfish analysis; llama.cpp local inference'],.8,2,11.8,25,.88)
call(s,'Source: notebooks/chess.ipynb and app/chess/stockfish_analyzer.py')
note(s,0,'این تنظیمات از دفترچه آموزش و کد پروژه خوانده شده‌اند. استفاده از متن بدون جفت موقعیت و توضیح، محدودیت اصلی آموزش است. جزئیات معماری عمومی ترنسفورمر از مسیر اصلی ارائه حذف شده‌اند تا زمان برای روش و نتیجه کافی باشد.')
s=slide('Benchmark provenance and interpretation','BACKUP / EVALUATION',True)
bullets(s,['Source games: Lichess, February 2013; at least 30 plies','Groups use mean player rating: <1200 / 1200–1799 / ≥1800','Pairing key: (type, row_id); 60 matched outputs per model','11 models; final saved outputs; fixed gemma4-12b judge'],.8,2,11.8,24,.92)
call(s,'Shared positions do not control hardware, model size, or generation settings.')
note(s,0,'این جزئیات برای سؤال درباره ساخت داده و قابلیت مقایسه ارائه شده‌اند. گروه ریتینگ توصیف بازی مبدأ است و نباید به عنوان برچسب مستقل دشواری تفسیر شود. تمام ۶۶۰ ردیف نهایی در محاسبه لحاظ شده‌اند.')
s=slide('Selected references and project evidence','BACKUP / REFERENCES',True)
refs=[('Jhamtani et al. (2018). Learning to Generate Move-by-Move Commentary.','https://aclanthology.org/P18-1154/'),('Zang et al. (2019). Automated Chess Commentator.','https://aclanthology.org/P19-1597/'),('Stockfish — chess analysis engine','https://stockfishchess.org/'),('llama.cpp — local model inference','https://github.com/ggml-org/llama.cpp')]
for i,(label,url) in enumerate(refs):
 y=1.9+i*.83;a=text(s,label,.8,y,11.8,.4,20);a.text_frame.paragraphs[0].runs[0].hyperlink.address=url;text(s,url,.8,y+.37,11.8,.3,14,GRAY)
text(s,'Project evidence: evaluation/evaluation_dataset.tsv; evaluation/judged__model_*.json;\nnotebooks/chess.ipynb; app/chess/stockfish_analyzer.py',.8,5.62,11.8,.9,17,TEAL)
note(s,0,'منابع منتخب برای مبنای پژوهش و ابزارهاست. اعداد نتایج از فایل‌های ارزیابی خود پروژه محاسبه شده‌اند. فهرست طولانی تاریخچه موتورهای شطرنج برای ارائه پانزده دقیقه‌ای ضروری نیست.')
prs.save(OUT/'BlunderGuard_15min.pptx')
assert sum(n[1] for n in notes)==855
with open(OUT/'Rehearsal_FA.md','w') as f:
 f.write('# راهنمای ارائه ۱۵ دقیقه‌ای\n\nاسلایدها انگلیسی و یادداشت‌های سخنران فارسی هستند. زمان گفتار پیشنهادی ۱۴ دقیقه و ۱۵ ثانیه است؛ ۴۵ ثانیه برای مکث و جابه‌جایی باقی می‌ماند. زمان پرسش و پاسخ جداست. اسلایدهای ۱۳ تا ۱۶ مخفی و مخصوص پرسش و پاسخ هستند. یادداشت‌ها راهنمای تمرین‌اند؛ زمان واقعی را با یک اجرای کامل بسنجید.\n\nنسخه نتایج نهایی: امتیاز مدل کوچک ۲٫۸۶ است و مقایسه آن با چند مدل بزرگ‌تر اضافه شده است. تاریخچه مفصل موتورهای شطرنج و جزئیات عمومی ترنسفورمر کوتاه شده‌اند؛ روش پروژه، داده، معیارها، نتایج و نمونه شکست اضافه شده‌اند. تصویر آغازین و نمودار مسیر از فایل اصلی حفظ شده‌اند. فایل اصلی تغییر نکرده است.\n\nمتن انگلیسی با Arial و یادداشت فارسی با B Nazanin تنظیم شده است. برای نمایش دقیق یادداشت‌ها فونت موجود در docs/report/fonts را روی دستگاه ارائه نصب کنید.\n\n')
 elapsed=0
 for i,(title,sec,t) in enumerate(notes,1):
  end=elapsed+sec
  f.write(f'## {i}. {title}\n\n'+(f'{elapsed//60:02}:{elapsed%60:02} تا {end//60:02}:{end%60:02} — {sec} ثانیه\n\n' if sec else 'پشتیبان — خارج از زمان اصلی\n\n')+t+'\n\n');elapsed=end
 f.write('## منابع داده و نام کامل مدل‌ها\n\nاعداد از evaluation/judged__model_*.json و خلاصه تأییدشده docs/report/chapter3_assets/verified_results.csv هستند؛ نتیجه نهایی مدل محلی از پرونده judged__model_chess-q8_0_repaired.json خوانده می‌شود.\n\n')
 for r,label in zip(rows,short):f.write(f'- {label}: `{r["model"]}`\n')
print(OUT/'BlunderGuard_15min.pptx')
